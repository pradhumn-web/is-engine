from __future__ import annotations

import re
from pathlib import Path
from typing import Any


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _fallback_pdf_text(raw: bytes) -> str:
    """Best-effort extraction from uncompressed PDF text operators."""
    chunks = re.findall(rb"\((?:\\.|[^\\)])*\)\s*Tj|\[(.*?)\]\s*TJ", raw, re.S)
    decoded: list[str] = []
    for chunk in chunks:
        if isinstance(chunk, tuple):
            chunk = chunk[0]
        data = chunk if isinstance(chunk, bytes) else b""
        parts = re.findall(rb"\((?:\\.|[^\\)])*\)", data)
        if not parts and data.startswith(b"("):
            parts = [data]
        for part in parts:
            value = part[1:-1]
            value = re.sub(rb"\\([\\()])", rb"\1", value)
            value = value.replace(rb"\\n", b" ").replace(rb"\\r", b" ")
            decoded.append(value.decode("latin-1", errors="ignore"))
    if not decoded:
        decoded = [m.decode("latin-1", errors="ignore") for m in re.findall(rb"[\x20-\x7e]{8,}", raw)]
    return "\n".join(decoded)


def parse_document(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = _decode_text(content)
    elif suffix == ".docx":
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        except Exception as exc:
            raise ValueError(f"Unable to parse DOCX: {exc}") from exc
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if len(text.strip()) < 20:
                text = _fallback_pdf_text(content)
        except Exception:
            text = _fallback_pdf_text(content)
        if not text.strip():
            raise ValueError("No extractable text found in PDF (scanned PDFs need OCR).")
    else:
        raise ValueError("Unsupported file type. Upload a .pdf, .docx, or .txt file.")
    return {"filename": filename, "text": text.strip(), "sections": segment_sections(text)}


def segment_sections(text: str) -> dict[str, str]:
    patterns = {
        "scope": r"(?im)^\s*(?:1[.)]?\s*)?(?:scope|work description|purpose)\s*[:\-]?\s*$",
        "technical_specs": r"(?im)^\s*(?:technical\s+specifications?|technical\s+requirements?|specifications?)\s*[:\-]?\s*$",
        "testing": r"(?im)^\s*(?:inspection\s*(?:and|&)\s*testing|testing|quality assurance|acceptance tests?)\s*[:\-]?\s*$",
        "compliance": r"(?im)^\s*(?:compliance|certification|standards|eligibility criteria)\s*[:\-]?\s*$",
    }
    matches: list[tuple[int, str, int]] = []
    for name, pattern in patterns.items():
        for m in re.finditer(pattern, text):
            matches.append((m.start(), name, m.end()))
    matches.sort()
    sections: dict[str, str] = {}
    for idx, (start, name, body_start) in enumerate(matches):
        end = matches[idx + 1][0] if idx + 1 < len(matches) else len(text)
        sections[name] = text[body_start:end].strip()
    sections["full_text"] = text.strip()
    return sections
